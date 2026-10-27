from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = Workbook()
ws = wb.active
ws.title = 'Leads'

headers = ['Business Name', 'Clean Website', 'Source', 'Priority', 'Facebook / Instagram Link', 'Contacted Date', 'Channel', 'Reply', 'Notes']

leads = [
    ['Julian Grant', 'facebook.com/Juliangrant1963', 'Reviews', 'Warm', 'facebook.com/Juliangrant1963', '', '', '', 'Already have FB link'],
    ['Mow and Grow Lawns', 'mowandgrowlawns.com', 'Facebook Ad', 'Hot', '', '', '', '', ''],
    ['CS Contractors', 'cscontractors.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Resin Driveways Quotes', 'resindrivewaysquotes.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Wrights Contractors', 'wrightscontractors.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Heritage Paving Leeds', 'heritagepavingleeds.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['The Green Horizon', 'thegreenhorizon.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Emily Joy Design', 'emilyjoydesign.com', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['FSL York', 'fslyork.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['DCM Surfaces', 'dcmsurfaces.com', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Norfolk Asphalt Driveways', 'norfolkasphaltdriveways.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['RM Associates', 'rmassoc.co.uk', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Proscapes Designs', 'proscapesdesigns.com', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['PW Groundworks', 'pwgroundworks.com', 'Sponsored', 'Hot', '', '', '', '', ''],
    ['Mr Landscape Garden', 'mrlandscapegarden.co.uk', 'Google Business', 'Warm', '', '', '', '', ''],
    ['Black Fox Garden Services', 'blackfoxgardenservices.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Helen Taylor Garden Design', 'helentaylorgardendesign.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Melissa Morton', 'melissamorton.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Thomas Moore Landscapes', 'thomasmoorelandscapes.com', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Flash Pyro', 'flashpyro.wixsite.com/website', 'Reviews', 'Warm', '', '', '', '', 'Very basic Wix site'],
    ['Elmwood Landscaping', 'elmwoodlandscaping.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Scorfield', 'scorfield.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['JD Landscapes and Ponds', 'jdlandscapesandponds.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Josh Snowden', 'joshsnowden.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Landmasons', 'landmasons.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Ellis Horticulture', 'ellishorticulture.com', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Sculpt Landscapes', 'sculptlandscapes.com', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Victoria Landscaping', 'victorialandscaping.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['D Sharp Landscapes', 'dsharplandscapes.com', 'Reviews', 'Warm', '', '', '', '', ''],
    ['MGS Landscapes', 'mgslandscapes.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['SR Miller', 'srmiller.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Paragon Services', 'paragonservices.co.uk', 'Reviews', 'Warm', '', '', '', '', ''],
    ['Roberto CC Landscape Inc', 'robertocclandscapeinc.com', 'Sponsored', 'CHECK', '', '', '', '', 'Sounds US - verify location first'],
    ['Holt Landscape Inc', 'holtlandscapeinc.com', 'Reviews', 'CHECK', '', '', '', '', 'Sounds US - verify location first'],
]

header_fill = PatternFill('solid', fgColor='1E1E2E')
hot_fill    = PatternFill('solid', fgColor='FFD6D6')
warm_fill   = PatternFill('solid', fgColor='D6F0FF')
alt_fill    = PatternFill('solid', fgColor='F7F7F7')
check_fill  = PatternFill('solid', fgColor='FFF3CD')

header_font = Font(bold=True, color='FFFFFF', size=11)
hot_font    = Font(bold=True, color='CC0000')
check_font  = Font(bold=True, color='996600')

thin   = Side(style='thin', color='CCCCCC')
border = Border(left=thin, right=thin, top=thin, bottom=thin)

for col, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border

ws.row_dimensions[1].height = 30

for row_idx, lead in enumerate(leads, 2):
    priority = lead[3]
    for col_idx, val in enumerate(lead, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.alignment = Alignment(vertical='center')
        cell.border = border
        if priority == 'Hot':
            cell.fill = hot_fill
            if col_idx == 4:
                cell.font = hot_font
        elif priority == 'CHECK':
            cell.fill = check_fill
            if col_idx == 4:
                cell.font = check_font
        elif row_idx % 2 == 0:
            cell.fill = warm_fill
        else:
            cell.fill = alt_fill
    ws.row_dimensions[row_idx].height = 18

col_widths = [26, 32, 16, 10, 34, 16, 18, 10, 30]
for i, w in enumerate(col_widths, 1):
    ws.column_dimensions[get_column_letter(i)].width = w

ws.freeze_panes = 'A2'
ws.auto_filter.ref = 'A1:I' + str(len(leads) + 1)

path = r'C:\Users\pirin\OneDrive\Mellékletek\personalized dm\outreach_tracker.xlsx'
wb.save(path)
print('Saved:', path)

"""
fill_niche.py — auto-fill the "Niche" column in a leads spreadsheet.

Uses the SAME classifier the DM app uses (en_niche), so the values match what
the app would detect. It only fills rows where the COMPANY NAME gives a
confident match; ambiguous/generic names are left BLANK on purpose so the app's
runtime Facebook-bio detection still runs for them at send time.

Usage:
    python fill_niche.py                       # uses the default path below
    python fill_niche.py "C:\\path\\to\\leads.xlsx"
"""
import sys
import openpyxl
from en_niche import detect_niche_from_name

DEFAULT_PATH = r"C:\Users\pirin\OneDrive\Asztali gép\uk trades.xlsx"


def fill(path: str) -> None:
    wb = openpyxl.load_workbook(path)
    ws = wb.active

    headers = [str(c.value or "").strip().lower() for c in ws[1]]

    # Locate the company-name column
    name_idx = next(
        (i for i, h in enumerate(headers) if "company" in h or "name" in h), None
    )
    if name_idx is None:
        raise SystemExit("Could not find a 'Company Name' column in row 1.")

    # Find an existing niche/trade column, or append a new "Niche" column
    niche_idx = next(
        (i for i, h in enumerate(headers) if "niche" in h or "trade" in h), None
    )
    if niche_idx is None:
        niche_idx = len(headers)
        ws.cell(row=1, column=niche_idx + 1, value="Niche")

    filled = blank = kept = 0
    for row in ws.iter_rows(min_row=2):
        name = str(row[name_idx].value or "").strip()
        if not name:
            continue

        cell = row[niche_idx] if niche_idx < len(row) else ws.cell(
            row=row[0].row, column=niche_idx + 1
        )

        # Don't clobber anything a human already typed
        if str(cell.value or "").strip():
            kept += 1
            continue

        niche = detect_niche_from_name(name)  # confident name-only match, else ""
        if niche:
            cell.value = niche
            filled += 1
            print(f"  {name[:45]:45s} -> {niche}")
        else:
            blank += 1  # left blank → app detects from FB bio at send time

    wb.save(path)
    print(f"\nDone. Filled {filled}, left {blank} blank for runtime detection, "
          f"kept {kept} existing.")
    print(f"Saved: {path}")


if __name__ == "__main__":
    fill(sys.argv[1] if len(sys.argv) > 1 else DEFAULT_PATH)

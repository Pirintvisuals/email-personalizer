"""
Auto DM Sender
──────────────
Run this script, then use the hotkey to work through your leads.

Hotkey: Ctrl+Shift+Space
  - 1st press → opens next lead in Chrome, generates DM, types it in the message box
  - You review it and press Enter to send
  - 2nd press → marks that lead as sent, waits a random delay, loads the next one

Daily cap: 30 sends (change DAILY_CAP below)
"""

import os, re, sys, time, random, threading
from pathlib import Path
import openpyxl
import pyperclip
import keyboard
from playwright.sync_api import sync_playwright

# ── Config ────────────────────────────────────────────────────────────────────
HOTKEY      = "ctrl+shift+space"
DAILY_CAP   = 30
MIN_DELAY   = 45   # seconds between leads (safety)
MAX_DELAY   = 90

# ── Load .env ─────────────────────────────────────────────────────────────────
_env = Path(__file__).parent / ".env"
if _env.exists():
    for line in _env.read_text().splitlines():
        if "=" in line and not line.startswith("#"):
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip())

# ── Import generation logic from dm_app ──────────────────────────────────────
from dm_app import (
    call_gemini, detect_trade,
    extract_prospect_location, extract_location_for_body,
    DM_BODY_ROOFING, DM_BODY_QUOTES,
    extract_first_name, patch_fallback_opener,
)


# ── Excel lead book ───────────────────────────────────────────────────────────
class LeadBook:
    def __init__(self, path: str):
        self.path = path
        self.wb   = openpyxl.load_workbook(path)
        self.ws   = self.wb.active
        headers   = [str(c.value or "").strip().lower() for c in self.ws[1]]

        self.url_col       = next((i for i, h in enumerate(headers) if "url" in h or "facebook" in h), 1)
        self.name_col      = next((i for i, h in enumerate(headers) if "company" in h or "name" in h), 0)
        self.city_col      = next((i for i, h in enumerate(headers) if "city" in h or "location" in h), None)
        self.niche_col     = next((i for i, h in enumerate(headers) if "niche" in h or "trade" in h), None)
        self.contacted_col = next((i for i, h in enumerate(headers) if "contact" in h), None)

    def next_lead(self):
        for row in self.ws.iter_rows(min_row=2):
            url       = str(row[self.url_col].value or "").strip()
            contacted = str(row[self.contacted_col].value or "").strip().lower() if self.contacted_col is not None else ""
            if url and contacted not in ("yes", "y", "sent", "1", "true"):
                return {
                    "row":   row[0].row,
                    "url":   url,
                    "name":  str(row[self.name_col].value or "").strip(),
                    "city":  str(row[self.city_col].value or "").strip()  if self.city_col  is not None else "",
                    "niche": str(row[self.niche_col].value or "").strip() if self.niche_col is not None else "",
                }
        return None

    def mark_sent(self, row_idx: int):
        if self.contacted_col is not None:
            self.ws.cell(row=row_idx, column=self.contacted_col + 1, value="Yes")
            self.wb.save(self.path)
            print(f"  Row {row_idx} marked as sent.")


# ── Browser / Facebook handler ────────────────────────────────────────────────
class Browser:
    def __init__(self):
        self.pw      = sync_playwright().start()
        self.browser = None
        self.ctx     = None
        self._connect()

    def _connect(self):
        try:
            self.browser = self.pw.chromium.connect_over_cdp("http://localhost:9222")
            self.ctx     = self.browser.contexts[0]
            print("Connected to Chrome.")
        except Exception as e:
            print(f"Chrome not found: {e}")
            print("Start Chrome with --remote-debugging-port=9222 first.")
            sys.exit(1)

    def open_url(self, url: str):
        page = self.ctx.new_page()
        try:
            page.goto(url, wait_until="domcontentloaded", timeout=45000)
        except Exception:
            # On timeout, try navigating with a lighter wait condition
            try:
                page.goto(url, wait_until="commit", timeout=30000)
            except Exception as e:
                raise RuntimeError(f"Could not load {url} — skipping. ({e})") from e
        page.wait_for_timeout(2500)
        return page

    def get_text(self, page) -> str:
        try:
            text = page.evaluate("""() => {
                ['script','style','noscript','svg','iframe']
                    .forEach(t => document.querySelectorAll(t).forEach(el => el.remove()));
                return document.body ? document.body.innerText : '';
            }""")
            text = re.sub(r"\n{3,}", "\n\n", text).strip()
            text = re.sub(r"(?im)^\s*[\d,\.]+\s*(followers?|likes?|reviews?|ratings?|check.?ins?|people follow|people like).*$", "", text)
            return re.sub(r"\n{3,}", "\n\n", text).strip()[:12000]
        except:
            return ""

    def click_message_button(self, page) -> bool:
        """Click the Send Message / Message button on a Facebook Page."""
        selectors = [
            '[aria-label="Send message"]',
            '[aria-label="Message"]',
            'a[href*="messenger.com/t/"]',
            'div[role="button"]:has-text("Message")',
            'a:has-text("Send message")',
            'a:has-text("Message")',
        ]
        for sel in selectors:
            try:
                el = page.locator(sel).first
                if el.is_visible(timeout=1500):
                    el.click()
                    page.wait_for_timeout(2500)
                    print(f"  Clicked message button ({sel})")
                    return True
            except:
                continue
        return False

    def type_into_message_box(self, dm: str) -> bool:
        """Find the Messenger input box across all open tabs and type the DM."""
        time.sleep(1.5)
        input_selectors = [
            '[aria-label="Message"]',
            '[aria-label="Aa"]',
            'div[contenteditable="true"][role="textbox"]',
            'div[contenteditable="true"]',
            'p.xat24cr',
        ]
        # Check newest tabs first (Messenger might have opened in a new tab)
        for pg in reversed(list(self.ctx.pages)):
            for sel in input_selectors:
                try:
                    el = pg.locator(sel).first
                    if el.is_visible(timeout=1500):
                        el.click()
                        # Human-like typing speed
                        for char in dm:
                            pg.keyboard.type(char)
                            time.sleep(random.uniform(0.02, 0.06))
                        print("  DM typed into message box.")
                        return True
                except:
                    continue
        return False


# ── Main controller ───────────────────────────────────────────────────────────
class AutoDM:
    def __init__(self, excel_path: str):
        self.book          = LeadBook(excel_path)
        self.browser       = Browser()
        self.sent_today    = 0
        self.current_lead  = None
        self.state         = "ready"  # ready | dm_typed
        self._lock         = threading.Lock()

    def hotkey_pressed(self):
        with self._lock:
            if self.state == "dm_typed":
                self._confirm_and_next()
            elif self.state == "ready":
                self._load_next()

    def _confirm_and_next(self):
        if self.current_lead:
            self.book.mark_sent(self.current_lead["row"])
            self.sent_today += 1
            print(f"\n  Sent today: {self.sent_today} / {DAILY_CAP}")

        if self.sent_today >= DAILY_CAP:
            print(f"\n  Daily cap of {DAILY_CAP} reached. Good work — come back tomorrow.")
            self.state = "ready"
            return

        delay = random.randint(MIN_DELAY, MAX_DELAY)
        print(f"  Waiting {delay}s before next lead...")
        time.sleep(delay)

        self.state = "ready"
        self._load_next()

    def _load_next(self):
        lead = self.book.next_lead()
        if not lead:
            print("\n  No more leads — all done!")
            return

        self.current_lead = lead
        print(f"\n{'=' * 55}")
        print(f"  Company : {lead['name']}")
        print(f"  URL     : {lead['url']}")
        print(f"  Opening profile in Chrome...")

        try:
            page      = self.browser.open_url(lead["url"])
            page_text = self.browser.get_text(page)

            # Build context for generation
            text_data = page_text
            if lead["niche"]:
                text_data = f"--- TRADE TYPE ---\n{lead['niche']}\n\n" + text_data
            if lead["city"]:
                text_data += f"\n\n--- LISTED CITY ---\n{lead['city']}"

            # Screenshot the posts area so the vision model sees actual job photos
            images = []
            try:
                page.evaluate("window.scrollTo(0, 850)")
                page.wait_for_timeout(1400)
                shot = page.screenshot(type="jpeg", quality=72)
                images = [(shot, "image/jpeg")]
                print(f"  Screenshot: {len(shot)} bytes")
                page.evaluate("window.scrollTo(0, 0)")
            except Exception as _se:
                print(f"  Screenshot skipped: {_se}")

            print("  Generating DM...")
            opener = call_gemini(text_data=text_data, images=images if images else None)
            opener = opener.rstrip(".!,") + "."
            opener = patch_fallback_opener(opener, text_data, lead.get("city", ""))

            trade = detect_trade(text_data)
            if trade == "roofers":
                location = extract_prospect_location("", text_data) or lead["city"] or "your area"
                body = DM_BODY_ROOFING.format(location=location)
            else:
                location = extract_location_for_body(opener, text_data)
                body = DM_BODY_QUOTES.format(location=location)

            dm = (opener + " " + body).strip() if body else opener
            print(f"\n  DM:\n  {dm}\n")

            # Copy to clipboard as backup
            pyperclip.copy(dm)

            # Click Message and type
            print("  Clicking Message button...")
            clicked = self.browser.click_message_button(page)
            if clicked:
                typed = self.browser.type_into_message_box(dm)
                if typed:
                    print("\n  Review the DM above, then press ENTER in Chrome to send.")
                    print("  After sending, press Ctrl+Shift+Space to mark sent and load next.\n")
                    self.state = "dm_typed"
                    return

            # Fallback — couldn't find message box
            print("\n  Could not find the message box automatically.")
            print("  DM copied to clipboard — click the message box and Ctrl+V to paste.")
            print("  After sending, press Ctrl+Shift+Space to mark sent and load next.\n")
            self.state = "dm_typed"

        except Exception as e:
            print(f"\n  Error: {e}")
            import traceback; traceback.print_exc()
            self.state = "ready"


# ── Entry point ───────────────────────────────────────────────────────────────
def find_excel() -> str:
    for f in Path(".").glob("*.xlsx"):
        name = f.name.lower()
        if any(k in name for k in ("lead", "dm", "contact", "prospect", "outreach")):
            return str(f)
    files = list(Path(".").glob("*.xlsx"))
    return str(files[0]) if files else ""


def main():
    print("\n  Auto DM Sender")
    print("  ──────────────")

    excel_path = find_excel()
    if not excel_path:
        excel_path = input("  Excel file path: ").strip().strip('"')
    if not os.path.isfile(excel_path):
        print(f"  File not found: {excel_path}")
        sys.exit(1)

    print(f"  Leads file : {Path(excel_path).name}")
    print(f"  Hotkey     : Ctrl+Shift+Space")
    print(f"  Daily cap  : {DAILY_CAP} sends")
    print(f"\n  READY. Press Ctrl+Shift+Space to load the first lead.\n")

    bot = AutoDM(excel_path)

    keyboard.add_hotkey(HOTKEY, lambda: threading.Thread(
        target=bot.hotkey_pressed, daemon=True
    ).start())

    keyboard.wait()


if __name__ == "__main__":
    main()
